# from openpyxl import Workbook, load_workbook
import xlrd
from docx import Document
root_path = "/home/ginger/Projects/Learning/P_Stack/file_tools/word_tool/20200416134916"
# wb = load_workbook(root_path + "/公司股本结构.xlsx")

# sheet = wb.get_sheet_by_name("sheet1")
# print(sheet)


def xlsx2docx(xlsx_path, docx_path):
    readbook = xlrd.open_workbook(xlsx_path)
    sheet = readbook.sheet_by_index(0)
    nrows = sheet.nrows
    ncols = sheet.ncols
    doc = Document()
    table = doc.add_table(rows=nrows, cols=ncols)
    table.style = "Medium Grid 1 Accent 1"
    for i in range(nrows):
        cur_row_values = sheet.row_values(i)
        cur_row_cells = table.rows[i].cells
        for j in range(ncols):
            cur_row_cells[j].text = str(cur_row_values[j])
    doc.save(docx_path)


xlsx_path = root_path + "/公司股本结构.xlsx"
docx_path = root_path + "/res.docx"
xlsx2docx(xlsx_path, docx_path)
